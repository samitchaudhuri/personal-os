#!/usr/bin/env python3
"""Write one center's site inputs into its ULC unit economics model.

Reads the site's row from combined_facts.csv, derives the model inputs (tier,
start date, rent, NNN, deposit, net leasehold), writes them into the center's
copy of the Baseline through Excel, has Excel recalculate and save, and reads
the cached results back with openpyxl.

Writes go through Excel via AppleScript, never openpyxl: openpyxl drops
drawings, threaded comments and dynamic-array metadata on save. openpyxl only
reads. The cell map lives in config.yaml under `unit_model`.
"""
from __future__ import annotations

import argparse
import csv
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import time
import warnings

import openpyxl
import yaml

from build_facts import parse_num

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.normpath(os.path.join(HERE, "..", ".."))
EXCEL = "Microsoft Excel"

# openpyxl only reads here, so the images it would drop on save don't matter.
warnings.filterwarnings("ignore", message="wmf image format", module="openpyxl")


def repo_path(p):
    """Relative config paths resolve against the personal-os repo root."""
    return p if os.path.isabs(p) else os.path.join(REPO, p)


def load_config(path=os.path.join(HERE, "config.yaml")):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def read_facts(cfg, site):
    """Return the combined_facts.csv row for `site` (exact display name)."""
    path = os.path.join(repo_path(cfg["drive_base"]), cfg["comparison_dir"], cfg["output_file"])
    with open(path, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["site"].strip() == site:
                return row
    raise SystemExit(f"no row for {site!r} in {path}")


def shift_month(year, month, offset):
    """(2027, 4), -3 -> (2027, 1); (2027, 1), -3 -> (2026, 10)."""
    i = year * 12 + (month - 1) + offset
    return i // 12, i % 12 + 1


def derive_inputs(row, cfg, start_offset=None, deposit=None, presales_rent=None):
    """Pure: map a combined_facts row to model input values, keyed like cfg inputs.

    rent and NNN use Python's round(), which rounds halves to even. SC Square's
    58 * 2247 / 12 is exactly 10860.5, and its existing copy holds 10860.
    """
    um = cfg["unit_model"]
    site = row["site"].strip()
    need = {k: parse_num(row.get(k)) for k in ("sf", "base_psf", "nnn_psf", "bo_net", "tier")}
    missing = [k for k, v in need.items() if v is None]
    m = re.fullmatch(r"(\d{4})-(\d{2})", (row.get("open_mo") or "").strip())
    if not m:
        missing.append("open_mo")
    if missing:
        raise ValueError(f"{site}: missing or TODO in combined_facts.csv: {', '.join(missing)}")

    offset = um["start_offset_months"] if start_offset is None else start_offset
    year, month = shift_month(int(m.group(1)), int(m.group(2)), offset)

    presales = presales_rent or um.get("presales_rent", {}).get(site)
    if presales not in ("zero", "full"):
        raise ValueError(f"{site}: set unit_model.presales_rent to zero or full in config.yaml")

    rent = round(need["base_psf"] * need["sf"] / 12)
    nnn = round(need["nnn_psf"] * need["sf"] / 12)
    return {
        "tier": f"Tier {int(need['tier'])}",
        "start_year": year,
        "start_month": month,
        "rent": rent,
        "rent_presales": rent if presales == "full" else 0,
        "nnn": nnn,
        "nnn_presales": nnn if presales == "full" else 0,
        "deposit": um["deposit"] if deposit is None else deposit,
        "net_leasehold": int(need["bo_net"]),
    }


def abatement_value(row, cfg, inputs):
    """Unmodeled upside: abated months x first-year all-in rent (0 if none)."""
    months = cfg["unit_model"].get("abatement_months", {}).get(row["site"].strip(), 0)
    return months * (inputs["rent"] + inputs["nnn"])


def cell_writes(inputs, cfg):
    """Pure: [(sheet, cell, value)] in cell-map order."""
    return [(sheet, cell, inputs[key])
            for key, targets in cfg["unit_model"]["inputs"].items()
            for sheet, cell in targets]


def formula_cells(path, writes):
    """Targets that hold a formula in the workbook; the model derives those."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return {(s, c) for s, c, _ in writes
                if isinstance(wb[s][c].value, str) and wb[s][c].value.startswith("=")}
    finally:
        wb.close()


def as_literal(v):
    """AppleScript literal for a cell value."""
    if isinstance(v, str):
        return '"' + v.replace("\\", "\\\\").replace('"', '\\"') + '"'
    return str(v)


def applescript(workbook_name, writes, settle_seconds=2):
    """Pure: script that writes the inputs under manual calc, recalcs and saves.

    Setting an input under automatic calc can crash Excel (AppleScript -609),
    so calc goes manual first. Excel may close the workbook itself after the
    save, hence the guarded close.
    """
    lines = [
        f'tell application "{EXCEL}"',
        f"\tset wb to workbook {as_literal(workbook_name)}",
        "\tset calculation to calculation manual",
    ]
    for sheet, cell, value in writes:
        lines.append(f"\tset value of range {as_literal(cell)} of worksheet "
                     f"{as_literal(sheet)} of wb to {as_literal(value)}")
    lines += [
        "\tset calculation to calculation automatic",
        "\tcalculate",
        f"\tdelay {settle_seconds}",
        "\tsave wb",
        "\ttry",
        "\t\tclose wb saving no",
        "\tend try",
        "end tell",
    ]
    return "\n".join(lines) + "\n"


def open_workbooks():
    out = subprocess.run(["osascript", "-e", f'tell application "{EXCEL}" to get name of workbooks'],
                         capture_output=True, text=True, check=True).stdout.strip()
    return [n.strip() for n in out.split(", ")] if out else []


def open_in_excel(path, timeout=90):
    """AppleScript `open workbook` on a Drive path silently fails; `open -a` works."""
    name = os.path.basename(path)
    if name in open_workbooks():
        raise SystemExit(f"{name} is already open in Excel; close it first")
    subprocess.run(["open", "-a", EXCEL, path], check=True)
    deadline = time.time() + timeout
    while time.time() < deadline:
        if name in open_workbooks():
            return name
        time.sleep(1)
    raise SystemExit(f"Excel did not open {name} within {timeout}s")


def run_excel(path, writes, settle_seconds=2):
    name = open_in_excel(path)
    time.sleep(settle_seconds)  # let the open finish its own recalc
    subprocess.run(["osascript", "-"], input=applescript(name, writes, settle_seconds),
                   text=True, check=True)


def read_results(path, cfg):
    """Cached values Excel saved: per-year EBITDA, peak cash, payback month."""
    res = cfg["unit_model"]["results"]
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[res["sheet"]]
    years = [ws[c].value for c in res["years"]]
    ebitda = [ws[c].value for c in res["ebitda"]]
    cash = [c.value for c in ws[res["cash_row"]] if isinstance(c.value, (int, float))]
    return {
        "ebitda": {int(y): round(v) for y, v in zip(years, ebitda)
                   if isinstance(y, (int, float)) and isinstance(v, (int, float))},
        "peak_cash": round(min(cash)) if cash else None,
        "payback_month": ws[res["payback"]].value,
    }


def check_writes(path, writes):
    """Every target's cached value must equal what was written (formula targets included)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    return [(s, c, v, wb[s][c].value) for s, c, v in writes if wb[s][c].value != v]


def ensure_copy(cfg, site, out=None, baseline=None, rebuild=False):
    """Return the center copy's path, creating it from the Baseline if missing."""
    um = cfg["unit_model"]
    models = repo_path(um["models_dir"])
    dest = out or os.path.join(models, um["copy_name"].format(site=site))
    src = baseline or os.path.join(models, um["baseline"])
    if rebuild and os.path.exists(dest):
        stem = os.path.splitext(os.path.basename(dest))[0]
        archived = os.path.join(models, um["archive_dir"],
                                f"{stem} {datetime.date.today():%Y-%m-%d}.xlsx")
        if os.path.exists(archived):
            raise SystemExit(f"{archived} already exists")
        shutil.move(dest, archived)
        print(f"archived {dest} -> {archived}")
    if not os.path.exists(dest):
        # -X: don't carry Drive's item-id xattr over to the copy.
        subprocess.run(["cp", "-X", src, dest], check=True)
        print(f"created {dest} from {os.path.basename(src)}")
    return dest


def main():
    ap = argparse.ArgumentParser(description="Write a site's inputs into its unit economics model")
    ap.add_argument("site", help='display name in combined_facts.csv, e.g. "SC Square"')
    ap.add_argument("--config", default=os.path.join(HERE, "config.yaml"))
    ap.add_argument("--out", help="model path (default: the center copy in models_dir)")
    ap.add_argument("--baseline", help="source for a missing copy (default: unit_model.baseline)")
    ap.add_argument("--rebuild", action="store_true", help="archive an existing copy, then copy the Baseline")
    ap.add_argument("--start-offset", type=int, help="months from open_mo to model start")
    ap.add_argument("--deposit", type=int, help="override the shared C146 deposit")
    ap.add_argument("--presales-rent", choices=["zero", "full"], help="override the site's presales rent")
    ap.add_argument("--dry-run", action="store_true", help="print the inputs and AppleScript; touch nothing")
    ap.add_argument("--json", action="store_true", help="print the results as JSON")
    args = ap.parse_args()

    cfg = load_config(args.config)
    row = read_facts(cfg, args.site)
    try:
        inputs = derive_inputs(row, cfg, args.start_offset, args.deposit, args.presales_rent)
    except ValueError as e:
        raise SystemExit(str(e))
    writes = cell_writes(inputs, cfg)

    if args.dry_run:
        print(json.dumps(inputs, indent=2))
        print(applescript(os.path.basename(args.out or cfg["unit_model"]["copy_name"].format(site=args.site)), writes))
        return

    path = ensure_copy(cfg, args.site, args.out, args.baseline, args.rebuild)
    skipped = formula_cells(path, writes)
    run_excel(path, [w for w in writes if (w[0], w[1]) not in skipped])

    bad = check_writes(path, writes)
    results = read_results(path, cfg)
    results["abatement_unmodeled"] = abatement_value(row, cfg, inputs)
    if args.json:
        print(json.dumps({"site": args.site, "inputs": inputs, "results": results}, indent=2))
    else:
        print(f"{args.site}: {path}")
        for k, v in inputs.items():
            print(f"  {k:<14} {v}")
        for s, c in sorted(skipped):
            print(f"  (left formula in {s}!{c})")
        for y, v in results["ebitda"].items():
            print(f"  EBITDA {y}    {v:,}")
        print(f"  peak cash      {results['peak_cash']:,}")
        print(f"  payback month  {results['payback_month']}")
        if results["abatement_unmodeled"]:
            print(f"  abatement      {results['abatement_unmodeled']:,} (not in the model)")
    if bad:
        for s, c, want, got in bad:
            print(f"ERROR: {s}!{c} is {got!r}, expected {want!r}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
