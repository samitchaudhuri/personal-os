"""Hermetic tests for the unit economics model runner (site_model.py).

These don't touch Google Drive or Excel. Derivations use today's SC Square and
Junction facts as golden values; the result reader runs on a small workbook
built in a temp dir.
"""
import copy
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import site_model as sm  # noqa: E402

CFG = sm.load_config()
SC_SQUARE = {"site": "SC Square", "sf": "2247", "base_psf": "58", "nnn_psf": "33.21",
             "bo_net": "146055", "tier": "4", "open_mo": "2027-04"}
JUNCTION = {"site": "Junction", "sf": "2399", "base_psf": "65", "nnn_psf": "21",
            "bo_net": "179925", "tier": "4", "open_mo": "2027-07"}


class TestShiftMonth(unittest.TestCase):
    def test_same_year(self):
        self.assertEqual(sm.shift_month(2027, 4, -3), (2027, 1))

    def test_wraps_into_prior_year(self):
        self.assertEqual(sm.shift_month(2027, 1, -3), (2026, 10))
        self.assertEqual(sm.shift_month(2027, 3, -3), (2026, 12))

    def test_zero_offset(self):
        self.assertEqual(sm.shift_month(2027, 7, 0), (2027, 7))


class TestDeriveInputs(unittest.TestCase):
    def test_sc_square(self):
        d = sm.derive_inputs(SC_SQUARE, CFG)
        self.assertEqual(d["tier"], "Tier 4")
        self.assertEqual((d["start_year"], d["start_month"]), (2027, 1))
        # 58 * 2247 / 12 = 10860.5 exactly; round() takes it to even, matching today's copy.
        self.assertEqual(d["rent"], 10860)
        self.assertEqual(d["nnn"], 6219)
        self.assertEqual(d["net_leasehold"], 146055)
        self.assertEqual(d["deposit"], 33000)

    def test_junction(self):
        d = sm.derive_inputs(JUNCTION, CFG)
        self.assertEqual((d["start_year"], d["start_month"]), (2027, 4))
        self.assertEqual((d["rent"], d["nnn"]), (12995, 4198))
        self.assertEqual(d["net_leasehold"], 179925)

    def test_presales_zero_by_default(self):
        d = sm.derive_inputs(SC_SQUARE, CFG)
        self.assertEqual((d["rent_presales"], d["nnn_presales"]), (0, 0))

    def test_regression_overrides(self):
        # How today's v.Jul3 copies were built: start = open_mo, full presales rent.
        d = sm.derive_inputs(SC_SQUARE, CFG, start_offset=0, deposit=46967, presales_rent="full")
        self.assertEqual((d["start_year"], d["start_month"]), (2027, 4))
        self.assertEqual((d["rent_presales"], d["nnn_presales"]), (10860, 6219))
        self.assertEqual(d["deposit"], 46967)

    def test_missing_facts_raise(self):
        row = dict(SC_SQUARE, nnn_psf="TODO", open_mo="")
        with self.assertRaisesRegex(ValueError, "nnn_psf.*open_mo"):
            sm.derive_inputs(row, CFG)

    def test_unset_presales_policy_raises(self):
        row = dict(SC_SQUARE, site="Camden Park")
        with self.assertRaisesRegex(ValueError, "presales_rent"):
            sm.derive_inputs(row, CFG)

    def test_abatement(self):
        d = sm.derive_inputs(JUNCTION, CFG)
        self.assertEqual(sm.abatement_value(JUNCTION, CFG, d), 12995 + 4198)
        self.assertEqual(sm.abatement_value(SC_SQUARE, CFG, sm.derive_inputs(SC_SQUARE, CFG)), 0)


class TestCellMap(unittest.TestCase):
    um = CFG["unit_model"]

    def test_every_input_has_a_target(self):
        self.assertEqual(set(self.um["inputs"]), set(sm.derive_inputs(SC_SQUARE, CFG)))

    def test_site_cells(self):
        mso = {k: [c for s, c in v if s == "(MSO) Inputs"] for k, v in self.um["inputs"].items()}
        self.assertEqual(mso["tier"], ["B29"])
        self.assertEqual((mso["start_year"], mso["start_month"]), (["C23"], ["D23"]))
        self.assertEqual((mso["rent"], mso["rent_presales"]), (["C129"], ["D129"]))
        self.assertEqual((mso["nnn"], mso["nnn_presales"]), (["C130"], ["D130"]))
        self.assertEqual((mso["deposit"], mso["net_leasehold"]), (["C146"], ["C147"]))

    def test_pc_start_date_mirrors_mso(self):
        # A PC/MSO start mismatch makes the PC Lifeforce-fee HLOOKUP return #N/A.
        for key in ("start_year", "start_month"):
            targets = self.um["inputs"][key]
            cells = {c for _, c in targets}
            self.assertEqual({s for s, _ in targets}, {"(MSO) Inputs", "(PC) Inputs"})
            self.assertEqual(len(cells), 1)

    def test_escalation_not_written(self):
        written = {c for v in self.um["inputs"].values() for s, c in v if s == "(MSO) Inputs"}
        self.assertNotIn("C131", written)

    def test_results(self):
        r = self.um["results"]
        self.assertEqual(r["sheet"], "(MSO) 3-Year IS")
        self.assertEqual(r["ebitda"], ["C118", "D118", "E118"])
        self.assertEqual((r["cash_row"], r["payback"]), (128, "I131"))

    def test_start_offset(self):
        self.assertEqual(self.um["start_offset_months"], -3)


class TestAppleScript(unittest.TestCase):
    def setUp(self):
        self.writes = sm.cell_writes(sm.derive_inputs(SC_SQUARE, CFG), CFG)
        self.script = sm.applescript("ULC Unit Economic Model - SC Square.xlsx", self.writes)
        self.lines = [l.strip() for l in self.script.splitlines()]

    def test_quoting(self):
        self.assertIn('set value of range "B29" of worksheet "(MSO) Inputs" of wb to "Tier 4"', self.lines)
        self.assertIn('set value of range "C129" of worksheet "(MSO) Inputs" of wb to 10860', self.lines)
        self.assertEqual(sm.as_literal('a "b" \\c'), '"a \\"b\\" \\\\c"')

    def test_manual_calc_around_writes(self):
        i = self.lines.index
        manual = i("set calculation to calculation manual")
        auto = i("set calculation to calculation automatic")
        sets = [n for n, l in enumerate(self.lines) if l.startswith("set value of range")]
        self.assertEqual(len(sets), len(self.writes))
        self.assertLess(manual, min(sets))
        self.assertGreater(auto, max(sets))
        self.assertLess(auto, i("calculate"))
        self.assertLess(i("calculate"), i("save wb"))

    def test_close_is_guarded(self):
        # Excel may close the workbook itself after saving.
        n = self.lines.index("close wb saving no")
        self.assertEqual(self.lines[n - 1], "try")

    def test_pc_writes(self):
        pc = [(c, v) for s, c, v in self.writes if s == "(PC) Inputs"]
        self.assertEqual(pc, [("B29", "Tier 4"), ("C23", 2027), ("D23", 1)])


class TestWorkbook(unittest.TestCase):
    """read_results, formula_cells and check_writes on a small built workbook."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "model.xlsx")
        wb = openpyxl.Workbook()
        inp = wb.active
        inp.title = "(MSO) Inputs"
        inp["B29"] = "Tier 4"
        inp["C129"] = 10860
        pc = wb.create_sheet("(PC) Inputs")
        pc["B29"] = "='(MSO) Inputs'!B29"
        is_ = wb.create_sheet("(MSO) 3-Year IS")
        for c, v in zip(("C8", "D8", "E8"), (2027, 2028, 2029)):
            is_[c] = v
        for c, v in zip(("C118", "D118", "E118"), (309716.4, 912996.3, 1034998.1)):
            is_[c] = v
        for col, v in zip("IJKLM", (-100.0, -982084.6, -500000.0, 12.5, 2.0)):
            is_[f"{col}128"] = v
        is_["H128"] = "label"
        is_["I131"] = 17
        wb.save(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_read_results(self):
        r = sm.read_results(self.path, CFG)
        self.assertEqual(r["ebitda"], {2027: 309716, 2028: 912996, 2029: 1034998})
        self.assertEqual(r["peak_cash"], -982085)
        self.assertEqual(r["payback_month"], 17)

    def test_formula_targets_are_skipped(self):
        writes = [("(MSO) Inputs", "B29", "Tier 4"), ("(PC) Inputs", "B29", "Tier 4")]
        self.assertEqual(sm.formula_cells(self.path, writes), {("(PC) Inputs", "B29")})

    def test_check_writes_flags_mismatch(self):
        writes = [("(MSO) Inputs", "B29", "Tier 4"), ("(MSO) Inputs", "C129", 10861)]
        self.assertEqual(sm.check_writes(self.path, writes),
                         [("(MSO) Inputs", "C129", 10861, 10860)])


class TestEnsureCopy(unittest.TestCase):
    def test_creates_missing_copy_and_keeps_existing(self):
        with tempfile.TemporaryDirectory() as d:
            cfg = copy.deepcopy(CFG)
            cfg["unit_model"]["models_dir"] = d
            with open(os.path.join(d, cfg["unit_model"]["baseline"]), "w") as f:
                f.write("baseline")
            dest = sm.ensure_copy(cfg, "SC Square")
            self.assertEqual(os.path.basename(dest), "ULC Unit Economic Model - SC Square.xlsx")
            with open(dest) as f:
                self.assertEqual(f.read(), "baseline")
            with open(dest, "w") as f:
                f.write("site edits")
            sm.ensure_copy(cfg, "SC Square")
            with open(dest) as f:
                self.assertEqual(f.read(), "site edits")


if __name__ == "__main__":
    unittest.main()
