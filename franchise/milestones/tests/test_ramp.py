import datetime as dt
import tempfile
import unittest
import zipfile
from unittest import mock
from pathlib import Path

import openpyxl
from openpyxl.chart import LineChart, Reference

import ramp
import timeline_csv as tc

D = dt.date


def lookup(r, col, sheet="Timeline"):
    c = "D" if col == "F" else "E"
    return (f'=IF(D{r}="Fixed",E{r},IFERROR(INDEX({sheet}!${c}$2:${c}$100,MATCH(D{r},{sheet}!$B$2:$B$100,0))+{col}{r},0))')


def make_ramp(path, sheet="Timeline"):
    wb = openpyxl.Workbook()
    wb.active.title = "Sites"
    s = wb["Sites"]
    s.append(["Site", "Name", "Lease sign", "Delivery", "Open"])
    s.append(["Site 1", "One"] + [f'=INDEX({sheet}!$D$2:$D$100,MATCH("{k}",{sheet}!$B$2:$B$100,0))'
                                  for k in ("lease-signed", "landlord-delivery", "soft-open")])
    d = wb.create_sheet("Draws")
    d.append(["Draw ID", "Site", "Purpose", "Anchor", "Fixed date", "Start", "End", "Start date", "End date"])
    d.append(["FEE", "Site 1", "Fee", "Fixed", D(2026, 6, 15), None, None, "=E2", "=E2"])
    d.append(["BO", "Site 1", "Build-out", "construction-start", None, 30, 30, lookup(3, "F", sheet), lookup(3, "G", sheet)])
    r = wb.create_sheet("README")
    r.append(["Sites tab: Site 1's dates look up the Timeline tab."])
    r.append(["Timeline tab: copied from the opening timeline CSV. Do not type in it."])
    ramp_ws = wb.create_sheet("Ramp")
    for row in ([1, 2, 3], [4, 5, 6]):
        ramp_ws.append(row)
    chart = LineChart()
    chart.add_data(Reference(ramp_ws, min_col=1, max_col=3, min_row=1, max_row=2), from_rows=True)
    ramp_ws.add_chart(chart, "E2")
    if sheet == "Timeline":
        wb.create_sheet("Timeline").append(["Id", "Key"])
    wb.save(path)


TABS = ramp.generated_tabs("junction")


class RampCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.xlsx = self.dir / "ramp.xlsx"
        make_ramp(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def moved(self):
        """Tabs as if the generator had moved construction-start a month later."""
        header, rows = TABS[ramp.MILESTONES]
        rows = [r[:3] + [r[3] + dt.timedelta(days=30), r[4]] + r[5:] if r[1] == "construction-start" else r for r in rows]
        return dict(TABS, **{ramp.MILESTONES: (header, rows)})


class TestTabs(unittest.TestCase):
    def test_key_start_and_end_sit_in_columns_b_d_and_e(self):
        header, rows = TABS[ramp.MILESTONES]
        self.assertEqual((header[1], header[3], header[4]), ("Key", "Start", "End"))
        by_key = {r[1]: r for r in rows}
        self.assertEqual(by_key["landlord-delivery"][3:6], [D(2027, 2, 8), D(2027, 2, 8), "assumed"])

    def test_milestones_include_the_presale_prerequisites(self):
        keys = [r[1] for r in TABS[ramp.MILESTONES][1]]
        for k in ("digital-partner-selected", "presale-kit-ordered", "membership-contract-reviewed"):
            self.assertIn(k, keys)

    def test_inputs_phases_and_rules_set_aside_come_from_the_generator(self):
        self.assertEqual({r[1] for r in TABS["Inputs"][1]} >= {"lease-signed", "landlord-delivery"}, True)
        self.assertIn("permit-review", [r[0] for r in TABS["Phases"][1]])
        self.assertEqual(TABS["Rules set aside"][1][0][0], "gm-in-seat to presale-launch")


class TestSync(RampCase):
    def test_writes_the_four_tabs_and_deletes_timeline(self):
        self.assertEqual(ramp.sync(self.xlsx, TABS), len(TABS[ramp.MILESTONES][1]))
        wb = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames[-4:], list(ramp.GENERATED))
        self.assertNotIn("Timeline", wb.sheetnames)

    def test_repoints_lookups_and_the_readme(self):
        ramp.sync(self.xlsx, TABS)
        wb = openpyxl.load_workbook(self.xlsx)
        self.assertIn("Milestones!$D$2:$D$100", wb["Draws"]["H3"].value)
        self.assertNotIn("Timeline", wb["Sites"]["C2"].value)
        self.assertEqual(wb["README"]["A1"].value, "Sites tab: Site 1's dates look up the Milestones tab.")
        self.assertTrue(wb["README"]["A2"].value.startswith("Milestones tab:"))

    def test_keeps_the_original_chart_bytes(self):
        before = zipfile.ZipFile(self.xlsx).read("xl/charts/chart1.xml")
        ramp.sync(self.xlsx, TABS)
        z = zipfile.ZipFile(self.xlsx)
        self.assertEqual(z.read("xl/charts/chart1.xml"), before)
        self.assertEqual(len(openpyxl.load_workbook(self.xlsx)["Ramp"]._charts), 1)

    def test_a_second_sync_replaces_the_tabs(self):
        ramp.sync(self.xlsx, TABS)
        ramp.sync(self.xlsx, self.moved())
        ws = openpyxl.load_workbook(self.xlsx)[ramp.MILESTONES]
        got = {r[1]: r[3].date() for r in ws.iter_rows(min_row=2, max_col=5, values_only=True) if r[1]}
        self.assertEqual(got["construction-start"], D(2027, 3, 10))

    def test_refuses_while_excel_has_the_file_open(self):
        (self.dir / "~$ramp.xlsx").write_text("")
        with self.assertRaises(SystemExit):
            ramp.sync(self.xlsx, TABS)


class TestCheck(RampCase):
    def test_passes_right_after_a_sync(self):
        ramp.sync(self.xlsx, TABS)
        problems, notes = ramp.check(self.xlsx, TABS)
        self.assertEqual(problems, [])
        self.assertIn("FEE", notes[0])

    def test_fails_when_the_generator_moves_after_the_sync(self):
        ramp.sync(self.xlsx, TABS)
        problems, _ = ramp.check(self.xlsx, self.moved())
        self.assertTrue(any("construction-start" in p and "Milestones tab differs" in p for p in problems))

    def test_fails_on_a_row_the_workbook_lacks_or_has_extra(self):
        ramp.sync(self.xlsx, TABS)
        header, rows = TABS[ramp.MILESTONES]
        problems, _ = ramp.check(self.xlsx, dict(TABS, **{ramp.MILESTONES: (header, rows + [[99, "new-key", "New", D(2027, 1, 1), D(2027, 1, 1), "computed", ""]])}))
        self.assertTrue(any("lacks new-key" in p for p in problems))
        problems, _ = ramp.check(self.xlsx, dict(TABS, **{ramp.MILESTONES: (header, rows[:-1])}))
        self.assertTrue(any("no longer writes" in p for p in problems))

    def test_fails_when_a_draw_names_a_key_the_generator_does_not_date(self):
        ramp.sync(self.xlsx, TABS)
        wb = openpyxl.load_workbook(self.xlsx)
        wb["Draws"]["D3"] = "permit-received"
        wb.save(self.xlsx)
        problems, _ = ramp.check(self.xlsx, TABS)
        self.assertTrue(any("permit-received" in p and "does not date" in p for p in problems))

    def test_fails_on_a_typed_draw_date_and_a_typed_site_date(self):
        ramp.sync(self.xlsx, TABS)
        wb = openpyxl.load_workbook(self.xlsx)
        wb["Draws"]["H3"] = D(2027, 3, 10)
        wb["Sites"]["E2"] = D(2027, 5, 31)
        wb.save(self.xlsx)
        problems, _ = ramp.check(self.xlsx, TABS)
        self.assertTrue(any("typed date in column H" in p for p in problems))
        self.assertTrue(any("Sites!E2" in p for p in problems))

    def test_fails_before_a_sync(self):
        problems, _ = ramp.check(self.xlsx, TABS)
        self.assertTrue(any("still has a Timeline tab" in p for p in problems))
        self.assertTrue(any("no Milestones tab" in p for p in problems))


class TestCsvHook(RampCase):
    def conf(self, csv_path):
        return {"timeline_csv": csv_path, "ramp": self.xlsx}

    def test_rewriting_the_published_csv_syncs_a_workbook_that_has_a_timeline_tab(self):
        csv_path = self.dir / "timeline.csv"
        with mock.patch.object(ramp, "outputs", return_value=self.conf(csv_path)):
            tc.sync_ramp("junction", csv_path)
        self.assertEqual(ramp.check(self.xlsx, TABS)[0], [])

    def test_a_csv_written_elsewhere_or_a_workbook_without_either_tab_is_left_alone(self):
        csv_path = self.dir / "timeline.csv"
        before = self.xlsx.read_bytes()
        with mock.patch.object(ramp, "outputs", return_value=self.conf(csv_path)):
            tc.sync_ramp("junction", self.dir / "other.csv")
        self.assertEqual(self.xlsx.read_bytes(), before)
        wb = openpyxl.load_workbook(self.xlsx)
        del wb["Timeline"]
        wb.save(self.xlsx)
        before = self.xlsx.read_bytes()
        with mock.patch.object(ramp, "outputs", return_value=self.conf(csv_path)):
            tc.sync_ramp("junction", csv_path)
        self.assertEqual(self.xlsx.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
