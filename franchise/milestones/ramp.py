"""Keep the funding ramp workbook on the generator's dates.

The ramp's Draws tab names a milestone key and day offsets for every draw the schedule dates, and a Milestones tab
holds the generator's dates for it to look up. `sync` writes four generated tabs, Inputs, Milestones, Phases and
Rules set aside, from a fresh generator run, and `check` fails when the workbook no longer agrees with one, so a
draw cannot sit on a date the schedule has moved. The ramp's path comes from `outputs` in sites.yaml.
"""
import argparse
import datetime as dt
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

import milestones as ms

MILESTONES = "Milestones"
OLD_TIMELINE = "Timeline"          # the tab this replaced; sync deletes it and repoints its lookups
GENERATED = ("Inputs", MILESTONES, "Phases", "Rules set aside")
LEGACY_ANCHORS = {"Fixed", "Lease sign", "Delivery", "Open"}
SITE_ANCHORS = {"C": "lease-signed", "D": "landlord-delivery", "E": "soft-open"}   # Sites tab columns on Site 1's row
DATE_FORMAT = "yyyy-mm-dd"
WIDTHS = {"Inputs": (6, 34, 40, 12, 12, 12, 70), MILESTONES: (6, 34, 40, 12, 12, 30, 70),
          "Phases": (28, 28, 28, 12, 12, 12, 12, 70), "Rules set aside": (44, 14, 70, 12)}
MILESTONES_NOTE = "Do not type here"
README_TEXT = ("Milestones tab: the generator's earliest and latest date for each Site 1 milestone, which the Draws and "
               "Sites tabs look up by Key. ramp.py sync writes it from a fresh generator run, and timeline_csv.py runs "
               "the sync whenever it rewrites the opening timeline CSV, which lists the ULC Los Gatos opening tasks "
               "for review and Gantt import and is not read here. Inputs, Phases and Rules set aside are generated "
               "beside it. Do not type in them. ramp.py check fails if a draw or any generated tab no longer matches "
               "a fresh generator run.")


def outputs(site):
    """The CSV and ramp paths for a site from sites.yaml, with ~ expanded."""
    conf = ms.load_yaml("sites.yaml")["sites"][site].get("outputs") or {}
    return {k: Path(v).expanduser() for k, v in conf.items()}


def as_date(value):
    return value.date() if isinstance(value, dt.datetime) else value if isinstance(value, dt.date) else None


def generated_tabs(site, durations=None, sites=None):
    """{tab: (header, rows)} from a fresh generator run. Key sits in column B of Inputs and Milestones, with the
    date or its start in D and its end in E, which is where the Draws lookups read."""
    durations = durations or ms.load_yaml("durations.yaml")
    sites = sites or ms.load_yaml("sites.yaml")
    conf = sites["sites"][site]
    solution = ms.solve(durations, conf["inputs"])
    resolved = solution.ranges
    return {
        "Inputs": (["Order", "Key", "Milestone", "Date from", "Date to", "Status", "Basis"],
                   [[n, k, label, lo, hi, status, basis]
                    for n, (k, label, (lo, hi), status, basis) in enumerate(ms.input_rows(durations, conf), 1)]),
        MILESTONES: (["Order", "Key", "Milestone", "Start", "End", "Status", "Basis"],
                     [[n, k, label, lo, hi, status, basis]
                      for n, (k, label, (lo, hi), status, _, basis)
                      in enumerate(ms.milestone_rows(durations, resolved, conf), 1)]),
        "Phases": (["Phase", "From", "To", "Planned min", "Planned max", "Actual min", "Actual max", "Basis"],
                   [[k, s, e, p[0], p[1], a[0], a[1], basis] for k, s, e, p, a, basis in ms.phase_rows(durations, resolved)]),
        "Rules set aside": (["Rule", "Kind", "Collides with", "Days short"], [list(r) for r in ms.set_aside_rows(solution.conflicts)]),
    }


def milestone_dates(tabs):
    """{key: (start, end)} from the Milestones tab rows."""
    return {r[1]: (r[3], r[4]) for r in tabs[MILESTONES][1]}


# ---- sync -------------------------------------------------------------------------------------------------

def repoint(wb):
    """Point every lookup and note that named the old Timeline tab at Milestones."""
    for name in ("Draws", "Sites"):
        if name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and f"{OLD_TIMELINE}!" in c.value:
                        c.value = c.value.replace(f"{OLD_TIMELINE}!", f"{MILESTONES}!")
    if "README" in wb.sheetnames:
        for row in wb["README"].iter_rows():
            for c in row:
                if isinstance(c.value, str) and OLD_TIMELINE in c.value:
                    c.value = README_TEXT if c.value.startswith(f"{OLD_TIMELINE} tab:") else c.value.replace(
                        f"{OLD_TIMELINE} tab", f"{MILESTONES} tab")
                elif isinstance(c.value, str) and c.value.startswith(f"{MILESTONES} tab:"):
                    c.value = README_TEXT


def sync(xlsx_path, tabs, recalc=False):
    """Write the generated tabs into the workbook, keeping its charts and everything else."""
    xlsx_path = Path(xlsx_path)
    lock = xlsx_path.with_name(f"~${xlsx_path.name}")
    if lock.exists():
        raise SystemExit(f"{xlsx_path.name} is open in Excel ({lock.name} exists); close it and rerun")
    wb = openpyxl.load_workbook(xlsx_path)
    for name in (OLD_TIMELINE, *GENERATED):
        if name in wb.sheetnames:
            del wb[name]
    for name in GENERATED:
        header, rows = tabs[name]
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.date):
                    c.number_format = DATE_FORMAT
        for i, width in enumerate(WIDTHS[name]):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = width
        ws.freeze_panes = "A2"
    ws = wb[MILESTONES]
    for n, (label, value) in enumerate([("Synced", dt.datetime.now().replace(microsecond=0)),
                                        (MILESTONES_NOTE, "ramp.py sync rewrites the generated tabs; the Draws tab "
                                                          "looks up Key, Start and End here")], 1):
        ws.cell(n, 9, label)
        ws.cell(n, 10, value)
    ws.column_dimensions["I"].width, ws.column_dimensions["J"].width = 18, 66
    repoint(wb)
    wb.calculation.fullCalcOnLoad = True
    originals = zipfile.ZipFile(xlsx_path)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False, dir=xlsx_path.parent) as tmp:
        tmp_path = Path(tmp.name)
    try:
        wb.save(tmp_path)
        restore_charts(originals, tmp_path)
        originals.close()
        os.replace(tmp_path, xlsx_path)
    finally:
        tmp_path.unlink(missing_ok=True)
    if recalc:
        recalculate(xlsx_path)
    return len(tabs[MILESTONES][1])


def restore_charts(original, new_path):
    """openpyxl rewrites charts and loses some of their formatting, so put the original chart parts back."""
    parts = {n: original.read(n) for n in original.namelist()
             if re.fullmatch(r"xl/(charts/[^/]+\.xml|drawings/drawing\d+\.xml|drawings/_rels/drawing\d+\.xml\.rels)", n)}
    if not parts:
        return
    with zipfile.ZipFile(new_path) as z:
        present = set(z.namelist())
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    missing = set(parts) - present
    if missing:
        raise SystemExit(f"the rewritten workbook lacks chart parts {sorted(missing)}; the original is untouched")
    with zipfile.ZipFile(new_path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, parts.get(info.filename, data))


def recalculate(xlsx_path):
    """Open the workbook in Excel so it recalculates and stores fresh values, then save and close it."""
    script = f'''tell application "Microsoft Excel"
  set wb to open workbook workbook file name POSIX file "{Path(xlsx_path).resolve()}"
  calculate full
  save wb
  close wb saving no
end tell'''
    done = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=180)
    if done.returncode:
        print(f"warning: Excel could not recalculate the ramp ({done.stderr.strip()}); it recalculates when opened",
              file=sys.stderr)


# ---- check ------------------------------------------------------------------------------------------------

def same(stored, expected):
    """Compare a stored cell with a generated value; Excel saves dates as datetimes and whole numbers as floats."""
    stored = as_date(stored) or stored
    return stored == expected or (isinstance(expected, int) and isinstance(stored, float) and stored == expected)


def check(xlsx_path, tabs):
    """Return (problems, notes) comparing the workbook with a fresh generator run's tabs."""
    wb = openpyxl.load_workbook(xlsx_path)
    cached = openpyxl.load_workbook(xlsx_path, data_only=True)
    problems, notes = [], []
    if OLD_TIMELINE in wb.sheetnames:
        problems.append(f"the workbook still has a {OLD_TIMELINE} tab; run ramp.py sync")
    for name in GENERATED:
        if name not in wb.sheetnames:
            problems.append(f"the workbook has no {name} tab; run ramp.py sync")
            continue
        header, rows = tabs[name]
        got = [r[:len(header)] for r in wb[name].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r[:len(header)])]
        if [c.value for c in wb[name][1]][:len(header)] != header:
            problems.append(f"{name} tab has different columns than the generator writes")
        for i in range(max(len(rows), len(got))):
            if i >= len(got):
                problems.append(f"{name} tab lacks {rows[i][1] if name in ('Inputs', MILESTONES) else rows[i][0]}")
            elif i >= len(rows):
                problems.append(f"{name} tab has {got[i][1] if name in ('Inputs', MILESTONES) else got[i][0]}, which the generator no longer writes")
            elif not all(same(a, b) for a, b in zip(got[i], rows[i])):
                label = rows[i][1] if name in ("Inputs", MILESTONES) else rows[i][0]
                problems.append(f"{name} tab differs from the generator at {label}: has {[as_date(v) or v for v in got[i][:5]]}, "
                                f"the generator gives {rows[i][:5]}")
    dates = milestone_dates(tabs)
    if MILESTONES not in wb.sheetnames:
        return problems, notes

    sites = wb["Sites"]
    for col in SITE_ANCHORS:
        v = sites[f"{col}2"].value
        if not (isinstance(v, str) and v.startswith("=") and f"{MILESTONES}!" in v):
            problems.append(f"Sites!{col}2 holds a typed {sites[f'{col}1'].value} date; it should look up {SITE_ANCHORS[col]}")

    draws, vals = wb["Draws"], cached["Draws"]
    fixed = []
    for r in range(2, draws.max_row + 1):
        ident, anchor = draws.cell(r, 1).value, draws.cell(r, 4).value
        if not ident or not anchor or str(ident).startswith("Amounts"):
            continue
        if anchor in LEGACY_ANCHORS:
            if anchor == "Fixed" and str(draws.cell(r, 2).value) == "Site 1":
                fixed.append(f"{ident} {as_date(draws.cell(r, 5).value)}")
            continue
        if anchor not in dates:
            problems.append(f"draw {ident} anchors to {anchor}, which the generator does not date")
            continue
        f, g = draws.cell(r, 6).value or 0, draws.cell(r, 7).value or 0
        start, end = dates[anchor]
        want = (start + dt.timedelta(days=f), end + dt.timedelta(days=g))
        for col, expect in ((8, want[0]), (9, want[1])):
            formula = draws.cell(r, col).value
            if not (isinstance(formula, str) and f"{MILESTONES}!" in formula):
                problems.append(f"draw {ident} has a typed date in column {'HI'[col - 8]}; it should look up {anchor}")
            got = as_date(vals.cell(r, col).value)
            if got is not None and got != expect:
                problems.append(f"draw {ident} shows {got} in column {'HI'[col - 8]} but {anchor} puts it at {expect}; "
                                "open the workbook in Excel to recalculate")
    if fixed:
        notes.append("Site 1 draws on typed dates, which no milestone moves: " + ", ".join(fixed))
    return problems, notes


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("command", choices=["sync", "check"])
    ap.add_argument("--site", default="junction")
    ap.add_argument("--ramp", help="ramp workbook; default is outputs.ramp in sites.yaml")
    ap.add_argument("--recalc", action="store_true", help="after a sync, recalculate the workbook in Excel")
    ap.add_argument("--skip-missing", action="store_true", help="with check, exit quietly when the workbook is absent, as when Drive is not mounted")
    args = ap.parse_args(argv)
    ramp = Path(args.ramp) if args.ramp else outputs(args.site).get("ramp")
    if not ramp:
        raise SystemExit(f"set outputs.ramp for {args.site} in sites.yaml, or pass --ramp")
    if args.command == "check" and args.skip_missing and not ramp.exists():
        print(f"skipped: {ramp} is not available")
        return
    tabs = generated_tabs(args.site)
    if args.command == "sync":
        print(f"synced {sync(ramp, tabs, args.recalc)} milestones from the {args.site} generator into {ramp.name}")
        return
    problems, notes = check(ramp, tabs)
    for n in notes:
        print(f"note: {n}")
    for p in problems:
        print(f"error: {p}", file=sys.stderr)
    if problems:
        raise SystemExit(f"{len(problems)} problem(s): the ramp does not follow the generator; run ramp.py sync")
    print(f"ok: {ramp.name} follows the {args.site} generator")


if __name__ == "__main__":
    main()
